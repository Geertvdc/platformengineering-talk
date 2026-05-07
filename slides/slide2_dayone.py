from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

SW, SH = 1920, 1080
W_IN, H_IN = 20.0, 11.25
EMU_PER_IN = 914400

def px(val, axis='x'):
    if axis == 'x':
        return int(val / SW * W_IN * EMU_PER_IN)
    else:
        return int(val / SH * H_IN * EMU_PER_IN)

def to_pt(pencil_px):
    return pencil_px * 72 / 96

PPT_PAD_LEFT = 9.6
PPT_PAD_TOP  = 4.8

prs = Presentation()
prs.slide_width  = Emu(int(W_IN * EMU_PER_IN))
prs.slide_height = Emu(int(H_IN * EMU_PER_IN))
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_rect(slide, x, y, w, h, hex_color, opacity=1.0):
    shape = slide.shapes.add_shape(1, px(x,'x'), px(y,'y'), px(w,'x'), px(h,'y'))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color.lstrip('#'))
    shape.line.fill.background()
    if opacity < 1.0:
        spF = shape.fill._xPr.find(qn('a:solidFill'))
        clr = spF.find(qn('a:srgbClr'))
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(opacity * 100000)))
    return shape

def add_text(slide, text, x, y, w, h,
             font_name, font_size_px, font_weight,
             hex_color, letter_spacing_px=0, extra_nudge_y=0, opacity=1.0):
    adj_x = x - PPT_PAD_LEFT
    adj_y = y - PPT_PAD_TOP - extra_nudge_y
    txBox = slide.shapes.add_textbox(px(adj_x,'x'), px(adj_y,'y'), px(w,'x'), px(h,'y'))
    tf = txBox.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    pPr = p._p.get_or_add_pPr()
    spcBef = etree.SubElement(pPr, qn('a:spcBef'))
    etree.SubElement(spcBef, qn('a:spcPts')).set('val', '0')
    spcAft = etree.SubElement(pPr, qn('a:spcAft'))
    etree.SubElement(spcAft, qn('a:spcPts')).set('val', '0')
    run = p.add_run()
    run.text = text
    rPr = run._r.get_or_add_rPr()
    rPr.set('sz', str(int(to_pt(font_size_px) * 100)))
    rPr.set('b', '1' if font_weight >= 700 else '0')
    if letter_spacing_px != 0:
        rPr.set('spc', str(int(to_pt(letter_spacing_px) * 100)))
    solidFill = etree.SubElement(rPr, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', hex_color.lstrip('#'))
    if opacity < 1.0:
        alpha = etree.SubElement(srgbClr, qn('a:alpha'))
        alpha.set('val', str(int(opacity * 100000)))
    weight_map = {400: '', 500: ' Medium', 600: ' SemiBold',
                  700: ' Bold', 800: ' ExtraBold', 900: ' Black'}
    latin = etree.SubElement(rPr, qn('a:latin'))
    latin.set('typeface', font_name + weight_map.get(font_weight, ''))
    return txBox

# ── Background (white) ───────────────────────────────────────────────────────
add_rect(slide, 0, 0, 1920, 1080, 'FFFFFF')

# ── Watermark "01" (x=1000, y=-120, fontSize=700, black, opacity=0.03) ───────
add_text(slide, '01',
         x=1000, y=-120, w=800, h=847,
         font_name='Inter', font_size_px=700, font_weight=900,
         hex_color='000000', letter_spacing_px=-30, opacity=0.03)

# ── Left accent bar (black on white) ─────────────────────────────────────────
add_rect(slide, 0, 0, 8, 1080, '000000')

# ── Tag row (content frame starts at x=120, y=100) ───────────────────────────
# Square: y=5 relative to tag frame → absolute y=105
add_rect(slide, 120, 105, 14, 14, '000000')
add_text(slide, '01 / THE CHAOS',
         x=150, y=100, w=500, h=24,
         font_name='Inter', font_size_px=20, font_weight=600,
         hex_color='000000', letter_spacing_px=4, extra_nudge_y=5)

# ── Title (y=253.5 relative to content frame → absolute y=353.5) ─────────────
# "DAY ONE." — fontSize=240, lineHeight=0.88 → h=211
add_text(slide, 'DAY ONE.',
         x=120, y=354, w=1600, h=211,
         font_name='Inter', font_size_px=240, font_weight=900,
         hex_color='000000', letter_spacing_px=-9)

# "Three things. Zero answers." — y=225 relative to title group → absolute y=354+225=579
add_text(slide, 'Three things. Zero answers.',
         x=120, y=579, w=800, h=36,
         font_name='Inter', font_size_px=30, font_weight=400,
         hex_color='AAAAAA', extra_nudge_y=3)

# ── 3 Columns (y=744 relative to content frame → absolute y=844) ─────────────
# Each column is 520px wide with gap=60. Col positions: 120, 700, 1280 (absolute x).
# Within each column (gap=10 between items):
#   Number text  y=0   relative → absolute y=844, h=87
#   Divider line y=97  relative → absolute y=941, h=2
#   Label text   y=109 relative → absolute y=953, h=27

col_x = [120, 700, 1280]
col_nums   = ['01', '02', '03']
col_labels = ['NAMESPACE.', 'DATABASE.', 'PIPELINE.']

for cx, num, label in zip(col_x, col_nums, col_labels):
    add_text(slide, num,
             x=cx, y=844, w=520, h=87,
             font_name='Inter', font_size_px=72, font_weight=900,
             hex_color='DDDDDD', letter_spacing_px=-3)
    add_rect(slide, cx, 941, 520, 2, '000000')
    add_text(slide, label,
             x=cx, y=953, w=520, h=27,
             font_name='Inter', font_size_px=22, font_weight=800,
             hex_color='000000', letter_spacing_px=2, extra_nudge_y=3)

out = '/Users/geertvdc/dev/github/geertvdc/platformengineering-talk/slides/slide2_dayone.pptx'
prs.save(out)
print('Saved:', out)
