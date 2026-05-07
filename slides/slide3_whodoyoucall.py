from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

SW, SH = 1920, 1080
W_IN, H_IN = 20.0, 11.25
EMU_PER_IN = 914400

def px(val, axis='x'):
    if axis == 'x':
        return int(val / SW * W_IN * EMU_PER_IN)
    else:
        return int(val / SH * H_IN * EMU_PER_IN)

def to_pt(pencil_px):
    return pencil_px * 72 / 96

PPT_PAD_LEFT = 9.6
PPT_PAD_TOP  = 4.8

prs = Presentation()
prs.slide_width  = Emu(int(W_IN * EMU_PER_IN))
prs.slide_height = Emu(int(H_IN * EMU_PER_IN))
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_rect(slide, x, y, w, h, hex_color, opacity=1.0):
    shape = slide.shapes.add_shape(1, px(x,'x'), px(y,'y'), px(w,'x'), px(h,'y'))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color.lstrip('#'))
    shape.line.fill.background()
    if opacity < 1.0:
        spF = shape.fill._xPr.find(qn('a:solidFill'))
        clr = spF.find(qn('a:srgbClr'))
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(opacity * 100000)))
    return shape

def add_oval(slide, x, y, w, h, hex_color, opacity=1.0):
    shape = slide.shapes.add_shape(9, px(x,'x'), px(y,'y'), px(w,'x'), px(h,'y'))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color.lstrip('#'))
    shape.line.fill.background()
    if opacity < 1.0:
        spF = shape.fill._xPr.find(qn('a:solidFill'))
        clr = spF.find(qn('a:srgbClr'))
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(opacity * 100000)))
    return shape

def add_text(slide, text, x, y, w, h,
             font_name, font_size_px, font_weight,
             hex_color, letter_spacing_px=0, extra_nudge_y=0, opacity=1.0):
    adj_x = x - PPT_PAD_LEFT
    adj_y = y - PPT_PAD_TOP - extra_nudge_y
    txBox = slide.shapes.add_textbox(px(adj_x,'x'), px(adj_y,'y'), px(w,'x'), px(h,'y'))
    tf = txBox.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    pPr = p._p.get_or_add_pPr()
    spcBef = etree.SubElement(pPr, qn('a:spcBef'))
    etree.SubElement(spcBef, qn('a:spcPts')).set('val', '0')
    spcAft = etree.SubElement(pPr, qn('a:spcAft'))
    etree.SubElement(spcAft, qn('a:spcPts')).set('val', '0')
    run = p.add_run()
    run.text = text
    rPr = run._r.get_or_add_rPr()
    rPr.set('sz', str(int(to_pt(font_size_px) * 100)))
    rPr.set('b', '1' if font_weight >= 700 else '0')
    if letter_spacing_px != 0:
        rPr.set('spc', str(int(to_pt(letter_spacing_px) * 100)))
    solidFill = etree.SubElement(rPr, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', hex_color.lstrip('#'))
    if opacity < 1.0:
        alpha = etree.SubElement(srgbClr, qn('a:alpha'))
        alpha.set('val', str(int(opacity * 100000)))
    weight_map = {400: '', 500: ' Medium', 600: ' SemiBold',
                  700: ' Bold', 800: ' ExtraBold', 900: ' Black'}
    latin = etree.SubElement(rPr, qn('a:latin'))
    latin.set('typeface', font_name + weight_map.get(font_weight, ''))
    return txBox

# ── Background (black) ───────────────────────────────────────────────────────
add_rect(slide, 0, 0, 1920, 1080, '000000')

# ── Watermark "?" (x=1280, y=-60, fontSize=700, white, opacity=0.05) ─────────
add_text(slide, '?',
         x=1280, y=-60, w=500, h=847,
         font_name='Inter', font_size_px=700, font_weight=900,
         hex_color='FFFFFF', opacity=0.05)

# ── RIGHT accent bar (x=1912) ────────────────────────────────────────────────
add_rect(slide, 1912, 0, 8, 1080, 'FFFFFF')

# ── Corner bar (bottom LEFT, x=0, y=1060, 200×8) ────────────────────────────
add_rect(slide, 0, 1060, 200, 8, 'FFFFFF')

# ── 3 dots on LEFT (single column at x=60, rows 160/200/240) ─────────────────
add_oval(slide, 60, 160, 16, 16, 'FFFFFF', 0.2)
add_oval(slide, 60, 200, 16, 16, 'FFFFFF', 0.4)
add_oval(slide, 60, 240, 16, 16, 'FFFFFF', 1.0)

# ── Tag row (content frame starts at x=120, y=120) ───────────────────────────
# Square: y=5 relative to tag frame → absolute y=125
add_rect(slide, 120, 125, 14, 14, 'FFFFFF')
add_text(slide, '01 / THE CHAOS',
         x=150, y=120, w=500, h=24,
         font_name='Inter', font_size_px=20, font_weight=600,
         hex_color='FFFFFF', letter_spacing_px=4, extra_nudge_y=5)

# ── Title (y=133.5 relative to content frame → absolute y=253.5) ─────────────
# 3 lines of 200px text, lineHeight=0.9 → step=180
add_text(slide, 'WHO DO',
         x=120, y=254, w=1600, h=180,
         font_name='Inter', font_size_px=200, font_weight=900,
         hex_color='FFFFFF', letter_spacing_px=-8)

add_text(slide, 'YOU',
         x=120, y=434, w=1600, h=180,
         font_name='Inter', font_size_px=200, font_weight=900,
         hex_color='666666', letter_spacing_px=-8)

add_text(slide, 'CALL?',
         x=120, y=614, w=1600, h=180,
         font_name='Inter', font_size_px=200, font_weight=900,
         hex_color='FFFFFF', letter_spacing_px=-8)

# ── Bottom (y=783 relative to content frame → absolute y=903) ────────────────
add_rect(slide, 120, 903, 120, 6, 'FFFFFF')
# Subtitle: y=26 relative to bottom group → absolute y=929
add_text(slide, 'When your team needs infrastructure. Today. Not next week.',
         x=120, y=929, w=1680, h=31,
         font_name='Inter', font_size_px=26, font_weight=400,
         hex_color='AAAAAA', extra_nudge_y=3)

out = '/Users/geertvdc/dev/github/geertvdc/platformengineering-talk/slides/slide3_whodoyoucall.pptx'
prs.save(out)
print('Saved:', out)
